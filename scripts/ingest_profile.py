"""이력서·경력기술서·포트폴리오를 profile-db에 적재한다 (텍스트 추출 → 청킹 → 임베딩 → 프로필 요약 갱신).

사용법:
  uv run python scripts/ingest_profile.py data/profile/resume.pdf --kind resume
  uv run python scripts/ingest_profile.py data/profile/*.md --kind portfolio
  uv run python scripts/ingest_profile.py --summary-only      # 문서 변경 없이 요약만 다시
지원 형식: .pdf .docx .md .txt  (같은 파일명을 다시 넣으면 내용이 바뀐 경우에만 교체)
"""
import argparse
import asyncio
import json
import logging
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mcp_server import server as db  # noqa: E402
from mcp_server.embeddings import embed_documents, text_hash, to_pg  # noqa: E402
from agents.llm import LLM  # noqa: E402
from agents.state import ProfileSummary  # noqa: E402

CHUNK_SIZE = 800
CHUNK_OVERLAP = 100
SUMMARY_MAX_CHARS = 60_000

logging.basicConfig(level=logging.WARNING)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif ext == ".docx":
        from docx import Document
        doc = Document(path)
        rows = ["\t".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
        text = "\n".join(p.text for p in doc.paragraphs) + ("\n\n" + "\n".join(rows) if rows else "")
    elif ext in (".md", ".txt"):
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError(f"지원하지 않는 형식: {path.name}")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk(text: str) -> list[str]:
    """문단 경계를 우선해 CHUNK_SIZE 내외로 자르고, 앞 조각 끝 CHUNK_OVERLAP자를 붙인다."""
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    pieces: list[str] = []
    for p in paras:   # 너무 긴 문단은 강제로 자른다
        pieces.extend(p[i:i + CHUNK_SIZE] for i in range(0, len(p), CHUNK_SIZE))
    chunks, cur = [], ""
    for p in pieces:
        if cur and len(cur) + len(p) + 2 > CHUNK_SIZE:
            chunks.append(cur)
            cur = cur[-CHUNK_OVERLAP:] + "\n\n" + p
        else:
            cur = f"{cur}\n\n{p}" if cur else p
    if cur:
        chunks.append(cur)
    return chunks


async def ingest(path: Path, kind: str) -> str:
    text = extract_text(path)
    if not text:
        return f"{path.name}: 추출된 텍스트 없음 (스캔 PDF라면 OCR 필요) — 건너뜀"
    h = text_hash(text)
    p = await db.pool()
    old = await p.fetchrow("SELECT id, content_hash FROM profile_documents WHERE filename = $1", path.name)
    if old and old["content_hash"] == h:
        return f"{path.name}: 변경 없음 — 건너뜀"
    chunks = chunk(text)
    vecs = await embed_documents(chunks)
    async with p.acquire() as conn, conn.transaction():
        doc_id = await conn.fetchval(
            """INSERT INTO profile_documents (kind, filename, content_hash, raw_text)
               VALUES ($1, $2, $3, $4)
               ON CONFLICT (filename) DO UPDATE
                 SET kind = EXCLUDED.kind, content_hash = EXCLUDED.content_hash,
                     raw_text = EXCLUDED.raw_text, updated_at = now()
               RETURNING id""",
            kind, path.name, h, text)
        await conn.execute("DELETE FROM profile_chunks WHERE document_id = $1", doc_id)
        await conn.executemany(
            "INSERT INTO profile_chunks (document_id, chunk_index, text, embedding) VALUES ($1,$2,$3,$4::vector)",
            [(doc_id, i, c, to_pg(v)) for i, (c, v) in enumerate(zip(chunks, vecs))])
    return f"{path.name}: {'교체' if old else '신규'} {len(text)}자 / {len(chunks)}조각"


async def summarize() -> dict | None:
    p = await db.pool()
    docs = await p.fetch("SELECT kind, filename, raw_text FROM profile_documents ORDER BY kind, filename")
    if not docs:
        await p.execute("DELETE FROM profile_summary")
        return None
    corpus = "\n\n".join(f"=== [{d['kind']}] {d['filename']} ===\n{d['raw_text']}" for d in docs)
    s = await LLM(schema=ProfileSummary, primary_timeout=60).ainvoke(
        "아래는 한 사람의 이력서·경력기술서·포트폴리오다. 문서에 적힌 사실만으로 구조화 프로필을 만들어라. "
        "추측하거나 과장하지 말고, 성과 수치는 문서 표현 그대로 옮긴다.\n\n" + corpus[:SUMMARY_MAX_CHARS])
    summary = s.model_dump()
    await p.execute(
        """INSERT INTO profile_summary (id, summary) VALUES (1, $1::jsonb)
           ON CONFLICT (id) DO UPDATE SET summary = EXCLUDED.summary, updated_at = now()""",
        json.dumps(summary, ensure_ascii=False))
    return summary


async def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="*", type=Path)
    ap.add_argument("--kind", choices=["resume", "career", "portfolio", "other"], default="other")
    ap.add_argument("--summary-only", action="store_true")
    args = ap.parse_args()
    if not args.files and not args.summary_only:
        ap.error("파일을 지정하거나 --summary-only를 쓰세요")
    changed = False
    for f in args.files:
        msg = await ingest(f, args.kind)
        changed |= "건너뜀" not in msg
        print(msg)
    if changed or args.summary_only:
        summary = await summarize()
        print("프로필 요약 갱신:", json.dumps(summary, ensure_ascii=False, indent=2) if summary else "(문서 없음)")


if __name__ == "__main__":
    asyncio.run(main())
